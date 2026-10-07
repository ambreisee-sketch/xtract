#!/usr/bin/env python3
"""
Junta à planilha de importação as matrículas resolvidas pelo conteúdo
(arquivos JSON em importacao_87/resolucoes/, um por matrícula, com as 7 abas já prontas).

Uso:
  python3 scripts/juntar_resolvidas.py --modelo "PLANILHA UNICA CONSOLIDADA - SCC 2209 MATRICULAS.xlsx" \\
      --base importacao_87/Matriculas_87_PADRAO_CONSOLIDADA.xlsx --resolucoes importacao_87/resolucoes \\
      --excluir 2745 --saida importacao_87/Matriculas_87_FINAL_PADRAO_CONSOLIDADA.xlsx
"""
import argparse, copy, glob, json, os, sys
sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.utils import get_column_letter
from gerar_planilha_importacao import ler_xlsx, chave_matricula, SHEETS
from validar_regras_importador import validar, agrupar

ap = argparse.ArgumentParser()
ap.add_argument('--modelo', required=True); ap.add_argument('--base', required=True)
ap.add_argument('--resolucoes', required=True); ap.add_argument('--excluir', nargs='*', default=[])
ap.add_argument('--saida', required=True)
a = ap.parse_args()

modelo = ler_xlsx(a.modelo)
headers = {}
for aba in SHEETS:
    h = modelo[aba][0][1]
    headers[aba] = [h[c][1] for c in sorted((c for c in h if h[c][1]), key=lambda c: (len(c), c))]

base = ler_xlsx(a.base)
linhas = {aba: [] for aba in SHEETS}
for aba in SHEETS:
    h = {cell[1]: c for c, cell in base[aba][0][1].items() if cell[1]}
    for _, r in base[aba][1:]:
        if chave_matricula(r.get('A')):
            linhas[aba].append([r[h[n]][1] for n in headers[aba]])
ja = {r[0] for r in linhas['MATRÍCULA']}

incluidas, fora = [], []
for f in sorted(glob.glob(os.path.join(a.resolucoes, '*.json')), key=lambda p: int(os.path.basename(p)[:-5])):
    rec = json.load(open(f, encoding='utf-8'))
    m = str(rec['matricula'])
    if m in a.excluir or rec.get('status') != 'resolvido' or m in ja:
        fora.append((m, 'excluída a pedido' if m in a.excluir else rec.get('status')))
        continue
    for aba in SHEETS:
        for row in rec['abas'].get(aba, []):
            linhas[aba].append([row[n] for n in headers[aba]])
    incluidas.append(m)

def chave(r):
    return int(r[0]) if str(r[0]).isdigit() else 10**9
for aba in SHEETS:
    linhas[aba].sort(key=chave)          # ordem por matrícula; ordem interna das linhas preservada

wb = openpyxl.load_workbook(a.modelo)
for aba in SHEETS:
    ws = wb[aba]; n = len(headers[aba])
    estilos = [copy.copy(ws.cell(row=2, column=j + 1)._style) for j in range(n)]
    altura = ws.row_dimensions[2].height
    ws.delete_rows(2, ws.max_row)
    for k in [k for k in ws.row_dimensions if k > 1]:
        del ws.row_dimensions[k]
    for i, vals in enumerate(linhas[aba]):
        for j, val in enumerate(vals):
            c = ws.cell(row=i + 2, column=j + 1, value=val); c.data_type = 's'; c._style = copy.copy(estilos[j])
        if altura:
            ws.row_dimensions[i + 2].height = altura
    ws.auto_filter.ref = f'A1:{get_column_letter(max(ws.max_column, n))}{len(linhas[aba]) + 1}'
wb.save(a.saida)
v = validar(agrupar(ler_xlsx(a.saida)))
print('base:', len(ja), '| resolvidas incluídas:', len(incluidas), '| fora:', fora)
print('total na planilha:', len({r[0] for r in linhas['MATRÍCULA']}), '| matrículas com problema nas regras do importador:', len(v))
