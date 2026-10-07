#!/usr/bin/env python3
"""
Monta uma planilha de importação, no padrão da consolidada, para uma LISTA de
matrículas, pegando cada matrícula da versão mais recente disponível.

  * As fontes são passadas em ordem de prioridade (a primeira que tiver a
    matrícula na aba MATRÍCULA é usada para todas as abas daquela matrícula).
  * Os valores passam pela mesma conversão do gerar_planilha_importacao.py
    (tudo texto, datas dd/mm/aaaa, "NÃO CONSTA" no vazio etc.).
  * Depois são aplicadas as regras do importador (scripts/validar_regras_importador.py):
      R3  PROTOCOLO do ato inválido -> "NÃO CONSTA" (o texto original continua no ato);
      R4  data do ato inexistente -> corrigida só se o próprio ato tiver a nota
          "onde se lê <data>, leia-se <data>"; senão vai para conferência;
      R2  parte com ID "NÃO CONSTA" -> ligada ao único ato com o mesmo nome de ato
          (ou cujo texto cita o nome da parte); se houver dúvida, vai para conferência;
      R1  numeração de ato repetida -> conferência (renumerar mudaria o livro).
  * DATA ABERTURA por extenso ("21 de novembro de 1985") vira 21/11/1985.
  * --renomear ANTIGA:NOVA@arquivo usa as linhas da matrícula ANTIGA daquele arquivo
    como se fossem da NOVA (para casos digitados com o número errado).
  * --conferir NUM="motivo" manda uma matrícula para conferência manual.

Gera <saida>/<nome>.xlsx (para importar), <saida>/<nome>_CONFERIR_NAO_IMPORTAR.xlsx
e <saida>/<nome>_relatorio.xlsx (origem de cada matrícula e cada ajuste).
"""
import argparse, collections, copy, datetime, os, re, sys

sys.dont_write_bytecode = True
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import openpyxl
from openpyxl.utils import get_column_letter
from gerar_planilha_importacao import (ler_xlsx, chave_matricula, converter, texto_data, norm_txt,
                                       norm_id, Log, NC, SHEETS)
from validar_regras_importador import validar, prot_ok, data_ok, agrupar
import gerar_planilha_importacao as _gpi
_gpi.ESTADO_COLS = set()   # as fontes já verificadas importam com o estado por extenso: não mexer

MESES = {'janeiro': 1, 'fevereiro': 2, 'marco': 3, 'abril': 4, 'maio': 5, 'junho': 6, 'julho': 7,
         'agosto': 8, 'setembro': 9, 'outubro': 10, 'novembro': 11, 'dezembro': 12}


def data_extenso(s):
    m = re.fullmatch(r'(\d{1,2})º?\s+de\s+([a-zç]+)\s+de\s+(\d{4})', norm_txt(s))
    if not m or m.group(2) not in MESES:
        return None
    try:
        return datetime.date(int(m.group(3)), MESES[m.group(2)], int(m.group(1))).strftime('%d/%m/%Y')
    except ValueError:
        return None


def errata(texto, data_atual):
    """'onde se lê 29 de Fevereiro de 2011, leia-se 29 de Fevereiro de 2012' -> '29/02/2012'."""
    t = norm_txt(texto)
    for m in re.finditer(r'onde se le,? (\d{1,2}º? de [a-zç]+ de \d{4}),? leia-se,? (\d{1,2}º? de [a-zç]+ de \d{4})', t):
        dm = re.fullmatch(r'(\d{1,2})º? de ([a-zç]+) de (\d{4})', m.group(1))
        if dm and dm.group(2) in MESES and '%02d/%02d/%s' % (int(dm.group(1)), MESES[dm.group(2)], dm.group(3)) == data_atual:
            return data_extenso(m.group(2))
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--modelo', required=True)
    ap.add_argument('--lista', required=True, help='arquivo texto/CSV com uma matrícula por linha')
    ap.add_argument('--fontes', nargs='+', required=True, help='em ordem de prioridade (mais recente primeiro)')
    ap.add_argument('--renomear', nargs='*', default=[], help='ANTIGA:NOVA@arquivo')
    ap.add_argument('--conferir', nargs='*', default=[], help='NUM=motivo')
    ap.add_argument('--saida', required=True)
    ap.add_argument('--nome', required=True)
    a = ap.parse_args()
    os.makedirs(a.saida, exist_ok=True)

    lista = [l.strip().lstrip('﻿') for l in open(a.lista, encoding='utf-8-sig') if l.strip()]
    lista = [x for x in dict.fromkeys(lista)]
    conferir_manual = dict(x.split('=', 1) for x in a.conferir)
    renome = {}
    for r in a.renomear:
        par, arq = r.split('@', 1)
        antiga, nova = par.split(':')
        renome[(os.path.basename(arq), antiga)] = nova

    modelo_raw = ler_xlsx(a.modelo)
    headers = {}
    for aba in SHEETS:
        h = modelo_raw[aba][0][1]
        cols = sorted((c for c in h if h[c][1]), key=lambda c: (len(c), c))
        headers[aba] = [h[c][1] for c in cols]

    # ---- linhas por matrícula, por fonte
    por_fonte = {}
    for f in a.fontes:
        nome = os.path.basename(f)
        dados = ler_xlsx(f)
        linhas = collections.defaultdict(lambda: collections.defaultdict(list))
        for aba in SHEETS:
            rows = dados[aba]
            h = {cell[1].strip(): c for c, cell in rows[0][1].items() if cell[1]}
            for rn, r in rows[1:]:
                k = chave_matricula(r.get('A'))
                if k is None:
                    continue
                k = renome.get((nome, k), k)
                if k not in lista:
                    continue
                linhas[k][aba].append(((nome, rn), {n: r.get(h[n]) for n in headers[aba] if n in h}))
        por_fonte[nome] = linhas

    log = Log()
    origem_de = {}
    ok_rows = {aba: [] for aba in SHEETS}
    conf = []                       # (mat, origem, motivo)
    conf_rows = {aba: [] for aba in SHEETS}
    ajustes = []                    # (mat, aba, origem, linha, coluna, antes, depois, motivo)

    for mat in lista:
        fonte = next((n for n in por_fonte if por_fonte[n][mat]['MATRÍCULA']), None)
        if fonte is None:
            conf.append((mat, '', 'não encontrada em nenhuma fonte'))
            continue
        L = por_fonte[fonte][mat]
        origem_de[mat] = fonte
        ctx = {}
        conv = {aba: [(o, [converter(aba, n, cells.get(n), mat, o, log, ctx) for n in headers[aba]])
                      for o, cells in L[aba]] for aba in SHEETS}
        for aba in SHEETS:          # a matrícula sai com o número da lista (caso renomeada)
            for o, v in conv[aba]:
                if v[0] != mat:
                    ajustes.append((mat, aba, o[0], o[1], 'MATRICULA', v[0], mat, 'matrícula digitada com o número errado na fonte'))
                    v[0] = mat

        def mandar_conferir(motivo):
            conf.append((mat, fonte, motivo))
            for aba in SHEETS:
                for o, v in conv[aba]:
                    conf_rows[aba].append([o[0], o[1]] + v)

        if mat in conferir_manual:
            mandar_conferir(conferir_manual[mat])
            continue
        HM = headers['MATRÍCULA']; iDA = HM.index('DATA ABERTURA')
        for o, v in conv['MATRÍCULA']:
            if not data_ok(v[iDA]) and v[iDA] != NC:
                d = data_extenso(v[iDA])
                if d:
                    ajustes.append((mat, 'MATRÍCULA', o[0], o[1], 'DATA ABERTURA', v[iDA], d, 'data por extenso -> dd/mm/aaaa'))
                    v[iDA] = d
        H = headers['MATRICULA_ATO']
        iP, iT, iO, iX, iD, iA = (H.index(x) for x in ('PROTOCOLO', 'TIPO ATO', 'ORDEM', 'TEXTO', 'DATA', 'ATO'))
        atos = conv['MATRICULA_ATO']
        ordens = [v[iO] for _, v in atos]
        if len(set(ordens)) != len(ordens):
            mandar_conferir('R1: numeração de ato repetida (%s)' % ' '.join('%s-%s' % (v[iT], v[iO]) for _, v in atos))
            continue
        problema = None
        for o, v in atos:
            if not prot_ok(v[iP]):
                ajustes.append((mat, 'MATRICULA_ATO', o[0], o[1], 'PROTOCOLO', v[iP], NC, 'R3: protocolo do ato inválido'))
                v[iP] = NC
            if not data_ok(v[iD]):
                nova = errata(v[iX], v[iD])
                if nova:
                    ajustes.append((mat, 'MATRICULA_ATO', o[0], o[1], 'DATA', v[iD], nova, 'R4: data inexistente corrigida pela nota do próprio ato'))
                    v[iD] = nova
                else:
                    problema = 'R4: data de ato inexistente (%s-%s %s)' % (v[iT], v[iO], v[iD])
        if problema:
            mandar_conferir(problema)
            continue
        ids = ['%s-%s' % (v[iT].strip().upper(), v[iO]) for _, v in atos]
        HP = headers['ATO_PARTES']; jI, jN, jP, jA = (HP.index(x) for x in ('ID', 'NOME DA PARTE', 'PROTOCOLO', 'ATO'))
        for o, v in conv['ATO_PARTES']:
            pid = norm_id(v[jI])
            if pid in ids:
                if v[jI] != pid:
                    ajustes.append((mat, 'ATO_PARTES', o[0], o[1], 'ID', v[jI], pid, 'ID padronizado'))
                    v[jI] = pid
                continue
            if v[jI] == NC:
                cand = [i for i, (_, x) in enumerate(atos) if norm_txt(x[iA]) == norm_txt(v[jA])]
                if len(cand) > 1:
                    nome = norm_txt(v[jN])
                    cand = [i for i in cand if nome and nome in norm_txt(atos[i][1][iX])] or cand
                if len(cand) == 1:
                    ajustes.append((mat, 'ATO_PARTES', o[0], o[1], 'ID', v[jI], ids[cand[0]], 'R2: parte sem ato ligada ao único ato com o mesmo nome'))
                    v[jI] = ids[cand[0]]
                    continue
            problema = 'R2: parte "%s" com ID %s sem ato correspondente' % (v[jN], v[jI])
            break
        if problema:
            mandar_conferir(problema)
            continue
        for o, v in conv['ATO_PARTES']:     # protocolo da parte acompanha o do ato
            i = ids.index(v[jI])
            if v[jP] != atos[i][1][iP]:
                ajustes.append((mat, 'ATO_PARTES', o[0], o[1], 'PROTOCOLO', v[jP], atos[i][1][iP], 'protocolo da parte = protocolo do ato'))
                v[jP] = atos[i][1][iP]
        nu, nr = len(conv['T. URBANO']), len(conv['T. RURAL'])
        if nu + nr != 1:
            mandar_conferir('imóvel sem descrição única (%d urbano, %d rural)' % (nu, nr))
            continue
        for aba in SHEETS:
            ok_rows[aba].extend(v for _, v in conv[aba])

    # ---- grava no modelo
    def gravar(dados, caminho):
        wb = openpyxl.load_workbook(a.modelo)
        for aba in SHEETS:
            ws = wb[aba]
            n = len(headers[aba])
            estilos = [copy.copy(ws.cell(row=2, column=j + 1)._style) for j in range(n)]
            altura = ws.row_dimensions[2].height
            ws.delete_rows(2, ws.max_row)
            for k in [k for k in ws.row_dimensions if k > 1]:
                del ws.row_dimensions[k]
            for i, vals in enumerate(dados[aba]):
                for j, val in enumerate(vals):
                    c = ws.cell(row=i + 2, column=j + 1, value=val)
                    c.data_type = 's'
                    c._style = copy.copy(estilos[j])
                if altura:
                    ws.row_dimensions[i + 2].height = altura
            ws.auto_filter.ref = f'A1:{get_column_letter(max(ws.max_column, n))}{len(dados[aba]) + 1}'
        wb.save(caminho)

    p_ok = os.path.join(a.saida, a.nome + '.xlsx')
    gravar(ok_rows, p_ok)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'LEIA'
    ws.append(['MATRICULA', 'VERSÃO USADA (fonte)', 'MOTIVO'])
    for r in conf: ws.append(list(r))
    for aba in SHEETS:
        w = wb.create_sheet(('CONF ' + aba)[:31]); w.append(['ARQUIVO DE ORIGEM', 'LINHA NA ORIGEM'] + headers[aba])
        for r in conf_rows[aba]: w.append(r)
    p_conf = os.path.join(a.saida, a.nome + '_CONFERIR_NAO_IMPORTAR.xlsx'); wb.save(p_conf)
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Origem'
    ws.append(['MATRICULA', 'VERSÃO USADA (fonte)', 'DESTINO'])
    conf_set = {c[0] for c in conf}
    for m in lista:
        ws.append([m, origem_de.get(m, ''), 'CONFERIR' if m in conf_set else 'IMPORTAR'])
    w = wb.create_sheet('Ajustes'); w.append(['MATRICULA', 'ABA', 'ARQUIVO', 'LINHA', 'COLUNA', 'ANTES', 'DEPOIS', 'MOTIVO'])
    for r in ajustes:
        if r[0] not in conf_set: w.append([str(x) for x in r])
    corrigidos = {(r[0], r[4]) for r in ajustes}
    for m_, aba, o0, o1, col, antes, depois, regra in log.ajustes:
        if m_ in conf_set or (regra.startswith('ATENÇÃO') and (m_, col) in corrigidos):
            continue
        w.append([m_, aba, o0, str(o1), col, str(antes), str(depois), regra])
    p_rel = os.path.join(a.saida, a.nome + '_relatorio.xlsx'); wb.save(p_rel)

    gravadas = len({r[0] for r in ok_rows['MATRÍCULA']})
    print('lista:', len(lista), '| importar:', gravadas, '| conferir:', len(conf))
    print('fontes usadas:', collections.Counter(origem_de.get(m, '-') for m in lista))
    v = validar(agrupar(ler_xlsx(p_ok)))
    print('validação das regras do importador na planilha gerada: matrículas com problema =', len(v))
    for p in (p_ok, p_conf, p_rel): print('gerado:', p)


if __name__ == '__main__':
    main()
