#!/usr/bin/env python3
"""
Gera planilhas de importação de matrículas no mesmo padrão da
"PLANILHA UNICA CONSOLIDADA - SCC 2209 MATRICULAS.xlsx" (que importa sem erro).

Regras aplicadas (copiadas do que a planilha consolidada faz):
  * todas as células gravadas como TEXTO (inline string), sem células numéricas;
  * datas como texto dd/mm/aaaa (DATA ABERTURA, DATA do ato, DATA DE NASCIMENTO);
  * AREA TOTAL numérica como texto com vírgula e 2 casas (ex.: 75,90);
  * célula vazia vira "NÃO CONSTA" (ex.: coluna NOTAS);
  * PROTOCOLO só com número (até 7 dígitos) ou número com ponto de milhar;
    qualquer outro conteúdo vira "NÃO CONSTA" (o original fica no relatório);
    o PROTOCOLO da parte (ATO_PARTES) acompanha o do ato a que ela pertence;
  * numeração dos atos (ORDEM) repetida dentro da matrícula é refeita em
    sequência única 1..n na ordem das linhas, e o ID das partes acompanha;
  * CNPJ/CPF/CEP que perderam zero à esquerda são completados;
  * linhas exatamente repetidas são removidas;
  * sem as ~344 mil células vazias "invisíveis" que existiam na planilha que falhou;
  * layout, estilos, larguras, filtro e painel congelado copiados da consolidada.

Matrículas que não podem ser corrigidas sem decisão humana (transcrições
divergentes da mesma matrícula, imóvel sem descrição urbana/rural ou com as
duas, parte que não se liga a nenhum ato) vão para uma planilha separada de
conferência e NÃO entram na planilha de importação.

Uso:
  python3 scripts/gerar_planilha_importacao.py \
      --modelo "PLANILHA UNICA CONSOLIDADA - SCC 2209 MATRICULAS.xlsx" \
      --situacao "Situacao_Matriculas_por_Livro (2).xlsx" \
      --ja-importadas "PLANILHA UNICA CONSOLIDADA - SCC 2209 MATRICULAS.xlsx" "LIVRO 2-AE (2).xlsx" \
      --fontes "Cópia de PLANILHA SANTA 2-AE a 2-AZ.xlsx" ... \
      --saida importacao_corrigida
"""
import argparse, collections, copy, datetime, os, re, sys, unicodedata, zipfile
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation
from xml.etree import ElementTree as ET

import openpyxl
from openpyxl.utils import get_column_letter

NC = 'NÃO CONSTA'
NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'
SHEETS = ['MATRÍCULA', 'T. URBANO', 'T. RURAL', 'MATRICULA_ATO', 'ATO_PARTES', 'P. FÍSICA', 'P. JURÍDICA']
DATE_COLS = {('MATRÍCULA', 'DATA ABERTURA'), ('MATRICULA_ATO', 'DATA'), ('P. FÍSICA', 'DATA DE NASCIMENTO')}
AREA_COLS = {('T. URBANO', 'AREA TOTAL'), ('T. RURAL', 'AREA TOTAL')}
PAD = {  # (aba, coluna): {tamanho_encontrado: tamanho_final}
    ('P. JURÍDICA', 'CNPJ/CGC'): {12: 14, 13: 14, 9: 11, 10: 11},
    ('P. JURÍDICA', 'CPF REPRESENTANTE 1'): {9: 11, 10: 11},
    ('P. JURÍDICA', 'CPF REPRESENTANTE 2'): {9: 11, 10: 11},
    ('P. FÍSICA', 'CPF'): {9: 11, 10: 11},
    ('P. FÍSICA', 'CPF DO CONJUGE'): {9: 11, 10: 11},
    ('ATO_PARTES', 'DOCUMENTO DA PARTE'): {9: 11, 10: 11, 12: 14, 13: 14},
    ('P. FÍSICA', 'CEP'): {7: 8}, ('P. JURÍDICA', 'CEP'): {7: 8},
    ('T. URBANO', 'CEP'): {7: 8}, ('T. RURAL', 'CEP'): {7: 8},
}
ILLEGAL = re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
EXCEL_EPOCH = datetime.date(1899, 12, 30)


# --------------------------------------------------------------------------- leitura
def ler_xlsx(path):
    """Lê o XML cru: {aba: [(linha, {coluna: (tipo, valor, formato)})]}."""
    z = zipfile.ZipFile(path)
    wb = ET.fromstring(z.read('xl/workbook.xml'))
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    ss = []
    if 'xl/sharedStrings.xml' in z.namelist():
        for si in ET.fromstring(z.read('xl/sharedStrings.xml')).findall('m:si', NS):
            ss.append(''.join(t.text or '' for t in si.iter('{%s}t' % NS['m'])))
    st = ET.fromstring(z.read('xl/styles.xml'))
    nf = st.find('m:numFmts', NS)
    numfmts = {int(n.get('numFmtId')): n.get('formatCode') for n in (nf if nf is not None else [])}
    xfs = [int(x.get('numFmtId', 0)) for x in st.find('m:cellXfs', NS)]
    out = {}
    for sh in wb.find('m:sheets', NS):
        tgt = rels[sh.get(RID)].lstrip('/')
        tgt = tgt if tgt.startswith('xl/') else 'xl/' + tgt
        rows = []
        for _, el in ET.iterparse(z.open(tgt)):
            if el.tag != '{%s}row' % NS['m']:
                continue
            r = {}
            for c in el.findall('m:c', NS):
                col = re.match(r'[A-Z]+', c.get('r')).group(0)
                t = c.get('t', 'n')
                v = c.find('m:v', NS)
                if t == 's':
                    val = ss[int(v.text)] if v is not None else None
                elif t == 'inlineStr':
                    isel = c.find('m:is', NS)
                    val = ''.join(x.text or '' for x in isel.iter('{%s}t' % NS['m'])) if isel is not None else None
                else:
                    val = v.text if v is not None else None
                s = int(c.get('s', 0))
                fid = xfs[s] if s < len(xfs) else 0
                r[col] = (t, val, numfmts.get(fid, 'builtin%d' % fid))
            rows.append((int(el.get('r')), r))
            el.clear()
        out[sh.get('name')] = rows
    return out


def chave_matricula(cell):
    if cell is None or cell[1] is None:
        return None
    s = str(cell[1]).strip()
    if not s:
        return None
    try:
        d = Decimal(s)
        if d == d.to_integral_value():
            return str(int(d))
    except InvalidOperation:
        pass
    s2 = s.replace('.', '')
    return s2 if s2.isdigit() else s


def norm_txt(s):
    s = unicodedata.normalize('NFKD', s or '')
    s = ''.join(ch for ch in s if not unicodedata.combining(ch))
    return re.sub(r'\s+', ' ', s).strip().lower()


# --------------------------------------------------------------------------- conversão
class Log:
    def __init__(self):
        self.ajustes = []      # (matricula, aba, arquivo, linha, coluna, antes, depois, regra)
        self.contagem = collections.Counter()

    def add(self, mat, aba, origem, col, antes, depois, regra, detalhar=True):
        self.contagem[regra] += 1
        if detalhar:
            self.ajustes.append((mat, aba, origem[0], origem[1], col, antes, depois, regra))


def serial_para_data(x):
    d = EXCEL_EPOCH + datetime.timedelta(days=int(x))
    return d.strftime('%d/%m/%Y')


def texto_data(s):
    m = re.fullmatch(r'\s*(\d{1,2})[/.-](\d{1,2})[/.-](\d{4})\s*', s)
    if not m:
        return None
    d, mo, y = map(int, m.groups())
    try:
        datetime.date(y, mo, d)
    except ValueError:
        return None
    return '%02d/%02d/%04d' % (d, mo, y)


def converter(aba, col, cell, mat, origem, log):
    """Converte uma célula de origem para o texto final."""
    t, v, fmt = cell if cell is not None else (None, None, None)
    if v is None or (isinstance(v, str) and v.strip() == ''):
        log.add(mat, aba, origem, col, '' if v is None else v, NC, 'vazio -> NÃO CONSTA', detalhar=False)
        return NC
    numerico = t not in ('s', 'inlineStr', 'str', 'e', 'b')
    if numerico:
        try:
            x = Decimal(str(v))
        except InvalidOperation:
            numerico = False
    if numerico:
        if (aba, col) in DATE_COLS and 1 <= x <= 2958465:
            out = serial_para_data(x)
            log.add(mat, aba, origem, col, str(v), out, 'data numérica -> texto dd/mm/aaaa', detalhar=False)
            return out
        if (aba, col) in AREA_COLS:
            out = str(x.quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)).replace('.', ',')
            regra = 'área numérica -> texto com vírgula'
            if 'd' in (fmt or '').lower() and 'mm' in (fmt or '').lower():
                regra = 'área com formato de data -> texto com vírgula'
                log.add(mat, aba, origem, col, str(v), out, regra)
            else:
                log.add(mat, aba, origem, col, str(v), out, regra, detalhar=False)
            return out
        if x == x.to_integral_value():
            out = str(int(x))
        else:
            out = format(x.normalize(), 'f').replace('.', ',')
            log.add(mat, aba, origem, col, str(v), out, 'número decimal -> texto com vírgula')
        if (aba, col) in PAD and out.isdigit() and len(out) in PAD[(aba, col)]:
            novo = out.zfill(PAD[(aba, col)][len(out)])
            log.add(mat, aba, origem, col, out, novo, 'zero à esquerda restaurado')
            out = novo
        return out
    s = str(v)
    s2 = ILLEGAL.sub('', s)
    if s2 != s:
        log.add(mat, aba, origem, col, repr(s), s2, 'caractere de controle removido')
    s = s2.strip()
    if s == '':
        return NC
    if norm_txt(s) in ('nao consta', 'n/c'):
        if s != NC:
            log.add(mat, aba, origem, col, s, NC, 'variação de NÃO CONSTA padronizada')
        return NC
    if (aba, col) in DATE_COLS:
        d = texto_data(s)
        if d and d != s:
            log.add(mat, aba, origem, col, s, d, 'data em texto padronizada dd/mm/aaaa')
            return d
        if not d:
            log.add(mat, aba, origem, col, s, s, 'ATENÇÃO: data em texto fora do padrão (mantida)')
    if (aba, col) in PAD and s.isdigit() and len(s) in PAD[(aba, col)]:
        novo = s.zfill(PAD[(aba, col)][len(s)])
        log.add(mat, aba, origem, col, s, novo, 'zero à esquerda restaurado')
        return novo
    if s.startswith('='):
        log.add(mat, aba, origem, col, s, s, 'texto começando com = gravado como texto')
    return s


def limpar_protocolo(p):
    if p == NC:
        return NC
    if re.fullmatch(r'\d{1,7}', p) or re.fullmatch(r'\d{1,3}(\.\d{3})+', p):
        return p
    return NC


def norm_id(s):
    m = re.fullmatch(r'\s*(R|AV)\s*[-.]?\s*0*(\d+)\s*', (s or '').upper())
    return '%s-%d' % (m.group(1), int(m.group(2))) if m else (s or '').strip().upper()


# --------------------------------------------------------------------------- montagem
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--modelo', required=True)
    ap.add_argument('--situacao', required=True)
    ap.add_argument('--ja-importadas', nargs='*', default=[])
    ap.add_argument('--fontes', nargs='+', required=True)
    ap.add_argument('--saida', required=True)
    a = ap.parse_args()
    os.makedirs(a.saida, exist_ok=True)

    # ---- quais matrículas faltam no sistema, por categoria do relatório
    ws = openpyxl.load_workbook(a.situacao, read_only=True, data_only=True)['Por Livro']
    hdr = None
    cat = {}
    livro_de = {}
    for row in ws.iter_rows(values_only=True):
        if hdr is None:
            hdr = list(row)
            continue
        if not row[0] or str(row[0]).startswith('TOTAL'):
            continue
        def nums(colname):
            v = row[hdr.index(colname)]
            return [int(x.replace('.', '')) for x in re.findall(r'\d[\d.]*', str(v))] if v else []
        for n in nums('DIGITADAS — quais são'):
            cat[str(n)] = 'DIGITADA'; livro_de[str(n)] = row[0]
        for n in nums('PLANILHA — quais são'):
            cat[str(n)] = 'PLANILHA'; livro_de[str(n)] = row[0]
        for n in nums('FALTANTES — quais são'):
            cat[str(n)] = 'FALTANTE'; livro_de[str(n)] = row[0]
    ja = set()
    for f in a.ja_importadas:
        for _, r in ler_xlsx(f)['MATRÍCULA'][1:]:
            k = chave_matricula(r.get('A'))
            if k:
                ja.add(k)
    faltam = {k: v for k, v in cat.items() if k not in ja}

    # ---- modelo (cabeçalhos e estilos)
    modelo_raw = ler_xlsx(a.modelo)
    headers = {}
    for aba in SHEETS:
        h = modelo_raw[aba][0][1]
        cols = sorted((c for c in h if h[c][1]), key=lambda c: (len(c), c))
        headers[aba] = [h[c][1] for c in cols]

    # ---- lê fontes e agrupa linhas por matrícula
    linhas = collections.defaultdict(lambda: collections.defaultdict(list))  # mat -> aba -> [(origem, {col: cell})]
    for f in a.fontes:
        dados = ler_xlsx(f)
        nome = os.path.basename(f)
        for aba in SHEETS:
            rows = dados[aba]
            h = {cell[1].strip(): c for c, cell in rows[0][1].items() if cell[1]}
            faltando = [n for n in headers[aba] if n not in h]
            if faltando:
                sys.exit(f'{nome} / {aba}: colunas ausentes {faltando}')
            for rn, r in rows[1:]:
                k = chave_matricula(r.get('A'))
                if k is None or k not in faltam:
                    continue
                linhas[k][aba].append(((nome, rn), {n: r.get(h[n]) for n in headers[aba]}))

    com_fonte = {k for k in linhas if linhas[k]['MATRÍCULA']}
    log = Log()
    saida = {cat_: {aba: [] for aba in SHEETS} for cat_ in ('importar', 'digitadas')}
    pendentes = []         # (mat, livro, categoria, motivo, detalhe)
    pend_linhas = {aba: [] for aba in SHEETS}
    atencao = []           # (mat, motivo, detalhe)
    renum = []             # (mat, linha origem, id antigo, id novo, data, ato)

    for mat in sorted(linhas, key=lambda x: int(x) if x.isdigit() else 10**9):
        L = linhas[mat]
        conv = {}
        for aba in SHEETS:
            conv[aba] = [(o, [converter(aba, n, cells[n], mat, o, log) for n in headers[aba]]) for o, cells in L[aba]]

        def pendente(motivo, detalhe=''):
            pendentes.append((mat, livro_de.get(mat, ''), faltam[mat], motivo, detalhe))
            for aba in SHEETS:
                for o, vals in conv[aba]:
                    pend_linhas[aba].append([o[0], o[1]] + vals)

        if not conv['MATRÍCULA']:
            continue  # matrícula só aparece em abas filhas: ignorada (não há cabeçalho da matrícula)

        # -- matrícula repetida (na mesma fonte ou em fontes diferentes)
        n_m = len(conv['MATRÍCULA'])
        if n_m > 1:
            iguais = len({tuple(v) for _, v in conv['MATRÍCULA']}) == 1
            filhos_ok = iguais and all(
                all(c % n_m == 0 for c in collections.Counter(tuple(v) for _, v in conv[aba]).values())
                for aba in SHEETS[1:])
            if not filhos_ok:
                origens = '; '.join(f'{o[0]} linha {o[1]}' for o, _ in conv['MATRÍCULA'])
                pendente('matrícula aparece mais de uma vez com dados diferentes', origens)
                continue
            for aba in SHEETS:
                vistos = set(); novo = []
                for o, v in conv[aba]:
                    if tuple(v) in vistos:
                        log.add(mat, aba, o, '(linha)', 'linha repetida', 'removida', 'cópia idêntica da matrícula removida')
                        continue
                    vistos.add(tuple(v)); novo.append((o, v))
                conv[aba] = novo

        # -- linhas exatamente repetidas nas abas filhas
        for aba in SHEETS[1:]:
            vistos = set(); novo = []
            for o, v in conv[aba]:
                if tuple(v) in vistos:
                    log.add(mat, aba, o, '(linha)', 'linha repetida', 'removida', 'linha idêntica removida')
                    continue
                vistos.add(tuple(v)); novo.append((o, v))
            conv[aba] = novo

        # -- exatamente uma descrição de imóvel (urbano OU rural)
        nu, nr = len(conv['T. URBANO']), len(conv['T. RURAL'])
        if nu + nr != 1:
            pendente('imóvel sem descrição única', f'{nu} linha(s) em T. URBANO e {nr} em T. RURAL')
            continue

        # -- atos: protocolo, numeração e ligação das partes
        H = headers['MATRICULA_ATO']; iP, iT, iO, iX, iD, iA = (H.index(x) for x in ('PROTOCOLO', 'TIPO ATO', 'ORDEM', 'TEXTO', 'DATA', 'ATO'))
        atos = conv['MATRICULA_ATO']
        ordens = [v[iO] for _, v in atos]
        if any(not o.isdigit() for o in ordens):
            pendente('ORDEM de ato não numérica', ', '.join(o for o in ordens if not o.isdigit()))
            continue
        for o, v in atos:
            p = limpar_protocolo(v[iP])
            if p != v[iP]:
                log.add(mat, 'MATRICULA_ATO', o, 'PROTOCOLO', v[iP], p, 'protocolo não numérico -> NÃO CONSTA')
                v[iP] = p
        antigos = ['%s-%d' % (v[iT].strip().upper(), int(v[iO])) for _, v in atos]
        precisa_renumerar = len(set(int(x) for x in ordens)) != len(ordens)
        if precisa_renumerar:
            datas = [datetime.datetime.strptime(v[iD], '%d/%m/%Y').date() for _, v in atos if texto_data(v[iD]) == v[iD]]
            if any(b < a for a, b in zip(datas, datas[1:])):
                pendente('numeração de atos repetida e atos fora de ordem cronológica',
                         ' '.join('%s-%s(%s)' % (v[iT], v[iO], v[iD]) for _, v in atos))
                continue
        novos = ['%s-%d' % (v[iT].strip().upper(), i + 1) for i, (_, v) in enumerate(atos)] if precisa_renumerar else antigos

        HP = headers['ATO_PARTES']; jI, jN, jP, jA = (HP.index(x) for x in ('ID', 'NOME DA PARTE', 'PROTOCOLO', 'ATO'))
        textos = [norm_txt(v[iX]) for _, v in atos]
        problema = None
        mapa_partes = []
        for o, v in conv['ATO_PARTES']:
            pid = norm_id(v[jI])
            cand = [i for i, x in enumerate(antigos) if x == pid]
            if len(cand) > 1:
                por_ato = [i for i in cand if norm_txt(atos[i][1][iA]) == norm_txt(v[jA])]
                cand = por_ato if len(por_ato) >= 1 else cand
            if len(cand) > 1:
                nome = norm_txt(v[jN])
                por_nome = [i for i in cand if nome and nome in textos[i]]
                cand = por_nome if len(por_nome) == 1 else cand
            if len(cand) != 1:
                problema = f'parte "{v[jN]}" com ID {v[jI]}: ' + ('nenhum ato com esse número' if not cand else 'mais de um ato possível')
                break
            mapa_partes.append(cand[0])
        if problema:
            pendente('parte não se liga a um único ato', problema)
            continue

        if precisa_renumerar:
            for i, (o, v) in enumerate(atos):
                if antigos[i] != novos[i]:
                    renum.append((mat, f'{o[0]} linha {o[1]}', antigos[i], novos[i], v[iD], v[iA]))
                    log.add(mat, 'MATRICULA_ATO', o, 'ORDEM', v[iO], str(i + 1), 'numeração de ato repetida -> sequência única', detalhar=False)
                v[iO] = str(i + 1)
        for (o, v), i in zip(conv['ATO_PARTES'], mapa_partes):
            if v[jI] != novos[i]:
                log.add(mat, 'ATO_PARTES', o, 'ID', v[jI], novos[i], 'ID da parte ajustado ao ato')
                v[jI] = novos[i]
            ap_ = atos[i][1][iP]
            if v[jP] != ap_:
                log.add(mat, 'ATO_PARTES', o, 'PROTOCOLO', v[jP], ap_, 'protocolo da parte = protocolo do ato')
                v[jP] = ap_
        # dedup de partes depois do ajuste
        vistos = set(); novo = []
        for o, v in conv['ATO_PARTES']:
            if tuple(v) in vistos:
                log.add(mat, 'ATO_PARTES', o, '(linha)', 'linha repetida', 'removida', 'linha idêntica removida')
                continue
            vistos.add(tuple(v)); novo.append((o, v))
        conv['ATO_PARTES'] = novo

        # -- avisos (entram na planilha, mas merecem conferência)
        dab = conv['MATRÍCULA'][0][1][headers['MATRÍCULA'].index('DATA ABERTURA')]
        if dab == NC:
            atencao.append((mat, 'DATA ABERTURA = NÃO CONSTA', 'a consolidada aceita, mas confira se a matrícula entrou'))
        if not atos:
            atencao.append((mat, 'matrícula sem atos', 'a consolidada tem casos assim'))
        elif not precisa_renumerar and sorted(int(x) for x in ordens) != list(range(1, len(ordens) + 1)):
            atencao.append((mat, 'numeração de atos com lacunas', ', '.join(antigos)))
        if precisa_renumerar:
            atencao.append((mat, 'atos renumerados', ' '.join(f'{x}->{y}' for x, y in zip(antigos, novos) if x != y)))

        destino = 'digitadas' if faltam[mat] == 'DIGITADA' else 'importar'
        for aba in SHEETS:
            saida[destino][aba].extend(v for _, v in conv[aba])

    # ---- grava planilhas de importação a partir do modelo
    def gravar(dados, caminho):
        wb = openpyxl.load_workbook(a.modelo)
        for aba in SHEETS:
            ws = wb[aba]
            ncols = len(headers[aba])
            estilos = [copy.copy(ws.cell(row=2, column=j + 1)._style) for j in range(ncols)]
            altura = ws.row_dimensions[2].height
            ws.delete_rows(2, ws.max_row)
            for i, vals in enumerate(dados[aba]):
                r = i + 2
                for j, val in enumerate(vals):
                    c = ws.cell(row=r, column=j + 1, value=val)
                    c.data_type = 's'
                    c._style = copy.copy(estilos[j])
                if altura:
                    ws.row_dimensions[r].height = altura
            ultima = max(ws.max_column, ncols)
            ws.auto_filter.ref = f'A1:{get_column_letter(ultima)}{len(dados[aba]) + 1}'
        wb.save(caminho)

    arquivos = []
    if any(saida['importar'][x] for x in SHEETS):
        p = os.path.join(a.saida, 'Matriculas_faltantes_PADRAO_CONSOLIDADA.xlsx'); gravar(saida['importar'], p); arquivos.append(p)
    if any(saida['digitadas'][x] for x in SHEETS):
        p = os.path.join(a.saida, 'Matriculas_DIGITADAS_encontradas_nas_copias_PADRAO_CONSOLIDADA.xlsx'); gravar(saida['digitadas'], p); arquivos.append(p)

    # ---- pendentes de conferência
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'LEIA'
    ws.append(['MATRICULA', 'LIVRO', 'CATEGORIA NO RELATÓRIO', 'MOTIVO', 'DETALHE'])
    for r in pendentes: ws.append(list(r))
    for aba in SHEETS:
        w = wb.create_sheet(aba)
        w.append(['ARQUIVO DE ORIGEM', 'LINHA NA ORIGEM'] + headers[aba])
        for r in pend_linhas[aba]: w.append(r)
    p = os.path.join(a.saida, 'PENDENTES_conferencia_manual_NAO_IMPORTAR.xlsx'); wb.save(p); arquivos.append(p)

    # ---- relatório
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Resumo'
    tot = lambda dest, aba: len(saida[dest][aba])
    ws.append(['Item', 'Quantidade'])
    ws.append(['Matrículas que faltam no sistema (relatório menos já importadas)', len(faltam)])
    ws.append(['Matrículas faltantes encontradas nas fontes', len(com_fonte)])
    ws.append(['Matrículas na planilha de importação', tot('importar', 'MATRÍCULA')])
    ws.append(['Matrículas DIGITADAS (planilha separada)', tot('digitadas', 'MATRÍCULA')])
    ws.append(['Matrículas pendentes de conferência manual', len(pendentes)])
    ws.append(['Matrículas faltantes sem nenhuma fonte', len(set(faltam) - com_fonte)])
    for aba in SHEETS:
        ws.append([f'Linhas em {aba} (importação)', tot('importar', aba)])
    ws.append([])
    ws.append(['Ajuste aplicado', 'Vezes'])
    for k, v in sorted(log.contagem.items(), key=lambda x: -x[1]): ws.append([k, v])
    w = wb.create_sheet('Ajustes'); w.append(['MATRICULA', 'ABA', 'ARQUIVO', 'LINHA', 'COLUNA', 'ANTES', 'DEPOIS', 'REGRA'])
    for r in log.ajustes: w.append([str(x) for x in r])
    w = wb.create_sheet('Atencao'); w.append(['MATRICULA', 'MOTIVO', 'DETALHE'])
    for r in atencao: w.append(list(r))
    w = wb.create_sheet('Atos_renumerados'); w.append(['MATRICULA', 'ORIGEM', 'ATO ANTES', 'ATO DEPOIS', 'DATA', 'ATO'])
    for r in renum: w.append(list(r))
    w = wb.create_sheet('Sem_fonte'); w.append(['MATRICULA', 'LIVRO', 'CATEGORIA NO RELATÓRIO'])
    for k in sorted(set(faltam) - com_fonte, key=int): w.append([k, livro_de.get(k, ''), faltam[k]])
    p = os.path.join(a.saida, 'Relatorio_ajustes.xlsx'); wb.save(p); arquivos.append(p)

    print('faltam no sistema:', len(faltam), '| com fonte:', len(com_fonte), '| importar:', tot('importar', 'MATRÍCULA'),
          '| digitadas:', tot('digitadas', 'MATRÍCULA'), '| pendentes:', len(pendentes), '| sem fonte:', len(set(faltam) - com_fonte))
    for k, v in sorted(log.contagem.items(), key=lambda x: -x[1]): print(f'   {v:7}  {k}')
    for p in arquivos: print('gerado:', p)


if __name__ == '__main__':
    main()
